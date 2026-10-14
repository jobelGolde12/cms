#!/usr/bin/env python3
"""Inspect the source workbook: sheets, dimensions, merged cells, formulas, headers, samples."""
import openpyxl

SRC = "School Form_ Personal Info of Students Needed.xlsx"

wb = openpyxl.load_workbook(SRC, data_only=False)
wb_data = openpyxl.load_workbook(SRC, data_only=True)

print("Sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    wsd = wb_data[name]
    print(f"\n===== SHEET: {name!r} =====")
    print(f"dim={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column} "
          f"state={ws.sheet_state}")
    # hidden rows/cols
    hidden_rows = [r for r, d in ws.row_dimensions.items() if d.hidden]
    hidden_cols = [c for c, d in ws.column_dimensions.items() if d.hidden]
    if hidden_rows: print(f"hidden_rows={hidden_rows}")
    if hidden_cols: print(f"hidden_cols={hidden_cols}")
    # merged cells
    merges = [str(m) for m in ws.merged_cells.ranges]
    print(f"merged_cells({len(merges)}): {merges[:25]}{'...' if len(merges)>25 else ''}")
    # formulas
    fcount = 0
    formulas = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                fcount += 1
                if len(formulas) < 8:
                    formulas.append(f"{c.coordinate}: {c.value}")
    print(f"formula_cells: {fcount} {formulas}")
    # first 12 rows preview (data_only values)
    print("--- first 12 rows (values) ---")
    for row in ws.iter_rows(min_row=1, max_row=min(12, ws.max_row)):
        vals = []
        for c in row[:15]:
            v = wsd.cell(row=c.row, column=c.column).value
            shown = "" if v is None else str(v)[:28]
            vals.append(shown)
        print(f"r{row[0].row}: {vals}")

# Full XML zip-level inspection for comments/defined names
import zipfile
z = zipfile.ZipFile(SRC)
print("\n===== ZIP CONTENTS (notable) =====")
for n in z.namelist():
    print(n)
